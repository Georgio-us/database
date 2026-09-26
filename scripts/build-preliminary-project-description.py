from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Предварительное описание проекта.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_font(run, size=None, bold=None, color="000000"):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def remove_style_borders(style):
    p_pr = style._element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_body(doc, text, bold_start=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.14
    if bold_start and text.startswith(bold_start):
        r1 = p.add_run(bold_start)
        configure_font(r1, 10.8, True)
        r2 = p.add_run(text[len(bold_start):])
        configure_font(r2, 10.8, False)
    else:
        run = p.add_run(text)
        configure_font(run, 10.8, False)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    configure_font(p.add_run(text), 10.6, False)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    configure_font(run, 15 if level == 1 else 12.2, True)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    configure_font(run, 9, False, "666666")


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.8)

    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
        remove_style_borders(style)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(9)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("Предварительное описание проекта")
    configure_font(title_run, 25, True)
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.space_after = Pt(22)
    subtitle_run = subtitle.add_run("Персональная база объектов агентства недвижимости")
    configure_font(subtitle_run, 13, False, "555555")

    add_body(doc, "Документ фиксирует предварительный состав проекта до подготовки технического задания. На его основании выбирается комплектация системы, согласуются границы разработки и определяется перечень функций для последующей детализации")
    add_body(doc, "После выбора комплектации готовится техническое задание с описанием полей, экранов, ролей, пользовательских действий, интеграций и критериев приёмки. Окончательный состав работ фиксируется после согласования технического задания")

    add_heading(doc, "1 Назначение системы")
    add_body(doc, "Система предназначена для ведения единой базы недвижимости агентства. В ней сотрудники смогут добавлять объекты, обновлять информацию, контролировать актуальность, назначать ответственных риелторов, формировать подборки для клиентов и публиковать объявления на внешних площадках")
    add_body(doc, "Работа выполняется через браузер на компьютере и телефоне. Основные пользователи - руководитель, администратор и риелторы агентства")
    add_body(doc, "Проект предусматривает три варианта комплектации стоимостью 1 800, 2 000 и 2 500 долларов")

    add_heading(doc, "2 База объектов стоимостью 1 800 долларов")
    add_body(doc, "В эту комплектацию входит полностью рабочая база, готовая к наполнению и ежедневному использованию")

    add_heading(doc, "2.1 Типы недвижимости", 2)
    for item in ("квартиры", "дома", "земельные участки", "коммерческая недвижимость", "паркинги"):
        add_bullet(doc, item)
    add_body(doc, "Система разрабатывается для продажи недвижимости. Аренда и посуточная аренда в текущий состав проекта не входят")

    sections = [
        ("2.2 Каталог объектов", [
            "Объекты отображаются в виде карточек и таблицы. В каталоге доступны поиск, фильтры, сортировка, выбор нескольких объектов и массовые действия",
            "Каждому объекту присваивается уникальный номер. Поиск работает по номеру объекта, адресу, району, микрорайону, типу недвижимости, цене, количеству комнат, площади, статусу и ответственному риелтору",
        ]),
        ("2.3 Карточка объекта", [
            "Карточка содержит основные сведения об объекте, включая тип недвижимости, адрес, район, микрорайон, цену, валюту, площадь, этаж, этажность, количество комнат, характеристики дома и помещения",
            "Набор полей меняется в зависимости от типа недвижимости. Для квартиры, дома, участка, коммерческого помещения и паркинга используются отдельные группы характеристик",
            "В карточке также хранятся фотографии, документы, внутренние комментарии, контакт собственника, ответственный риелтор и сведения об эксклюзивности объекта",
        ]),
        ("2.4 Адреса и районы", [
            "База ориентирована на работу с недвижимостью Одессы. В системе используются четыре административных района, микрорайоны и дополнительные территориальные обозначения",
            "Точный адрес хранится внутри базы. В клиентских подборках номер квартиры, контакт собственника и внутренние комментарии не показываются",
        ]),
        ("2.5 Контакты и ответственный риелтор", [
            "Для собственника можно сохранить имя, несколько телефонных номеров, Telegram, Viber, WhatsApp и внутренний комментарий",
            "За объектом закрепляется ответственный риелтор. Другие сотрудники могут подтвердить актуальность, изменить цену или статус, добавить комментарий и зафиксировать результат разговора с собственником. Все изменения сохраняются в журнале",
            "Смена ответственного риелтора доступна руководителю и администратору",
        ]),
        ("2.6 Статусы и актуализация", [
            "Для объектов используются статусы черновик, активный, резерв и архив. При переносе в архив указывается причина, например продажа или снятие с продажи",
            "Дата создания, дата последнего изменения и дата подтверждения актуальности хранятся отдельно",
            "Система напоминает о необходимости проверить объект через 7 или 14 дней. Просроченная актуализация не удаляет объект и не снимает его с публикации автоматически",
        ]),
        ("2.7 Роли и права доступа", [
            "Предусмотрены роли руководителя, администратора и риелтора",
            "Все сотрудники видят общую базу и контакты собственников. Права на изменение отдельных данных, управление пользователями, смену ответственных и системные настройки определяются ролью",
        ]),
        ("2.8 История изменений", [
            "В журнале фиксируются изменение цены, статуса, ответственного риелтора, контактных данных, характеристик и актуальности объекта",
            "Для каждой записи сохраняются сотрудник, дата, время и содержание изменения",
        ]),
        ("2.9 Фотографии документы и эксклюзивы", [
            "Для объекта можно загрузить фотографии и изменить порядок их отображения. Отдельные фотографии можно скрыть из клиентской подборки",
            "Документы загружаются при необходимости и остаются доступны сотрудникам агентства",
            "Эксклюзивный объект получает отдельную отметку и дату начала эксклюзива",
        ]),
        ("2.10 Импорт и экспорт", [
            "Объекты можно загружать из подготовленных файлов Excel или CSV. Перед сохранением система показывает найденные ошибки и позволяет проверить сопоставление столбцов",
            "Данные из каталога можно выгружать в Excel или CSV для дальнейшей работы",
        ]),
        ("2.11 Клиентские подборки", [
            "Риелтор выбирает один или несколько объектов и формирует веб-подборку. Подборка открывается по отдельной ссылке и показывает только разрешённую для клиента информацию",
            "Ссылку можно скопировать или отправить через Telegram. В базе фиксируется открытие подборки клиентом",
            "Создание отдельного Telegram-бота и подключение личных аккаунтов сотрудников в эту комплектацию не входят",
        ]),
        ("2.12 Валюты", [
            "Цена объекта указывается в долларах или гривне. Для пересчёта используется курс Национального банка Украины, который обновляется один раз в день",
            "Исходная цена собственника сохраняется без изменения. Пересчитанная сумма используется для отображения",
        ]),
        ("2.13 Публикация на OLX и DIM RIA", [
            "Объекты можно отправлять на OLX и DIM RIA через один корпоративный аккаунт агентства для каждой площадки",
            "В системе отображаются состояние публикации, дата отправки и сообщение об ошибке. Формат и доступные действия зависят от технических возможностей аккаунтов и условий подключения площадок",
            "Доступы к корпоративным кабинетам и необходимые разрешения предоставляются со стороны агентства",
        ]),
        ("2.14 Работа с телефона", [
            "Основные действия доступны с телефона, включая поиск, просмотр карточки, создание объекта, редактирование данных, актуализацию, работу с фотографиями и формирование подборки",
        ]),
        ("2.15 Выгрузка на сайт", [
            "Выгрузка объектов на сайт в стоимость 1 800 долларов не входит, поскольку сайт на данный момент отсутствует и его техническая структура неизвестна",
            "Интеграция проектируется после появления сайта. Тогда определяется формат обмена данными, состав публикуемых полей, правила обновления и порядок удаления объектов",
        ]),
    ]

    for heading, paragraphs in sections:
        add_heading(doc, heading, 2)
        for paragraph in paragraphs:
            add_body(doc, paragraph)

    add_heading(doc, "3 База объектов и заявки стоимостью 2 000 долларов")
    add_body(doc, "Комплектация включает весь функционал базы стоимостью 1 800 долларов и отдельный раздел для работы с заявками клиентов")

    add_heading(doc, "3.1 Что такое заявка", 2)
    add_body(doc, "Заявка фиксирует запрос клиента на недвижимость. В ней указываются контакт клиента, ответственный риелтор, тип недвижимости, районы, количество комнат, площадь, бюджет и дополнительные пожелания")
    add_body(doc, "Пример заявки - двухкомнатная квартира в Приморском районе площадью от 55 квадратных метров стоимостью до 100 000 долларов")

    add_heading(doc, "3.2 Работа с заявкой", 2)
    add_body(doc, "Система показывает объекты из текущей базы, которые соответствуют указанным параметрам. Риелтор может добавить в заявку любой объект вручную")
    add_body(doc, "Объекты внутри заявки получают отметки новый, подходит, не подходит, предложен клиенту и просмотрен. Из подходящих вариантов формируется клиентская веб-подборка")
    add_body(doc, "При добавлении нового объекта система проверяет его соответствие активным заявкам. Подходящие варианты появляются внутри соответствующих заявок")
    add_body(doc, "Заявка может быть активной, приостановленной или закрытой")

    add_heading(doc, "3.3 Границы модуля заявок", 2)
    add_body(doc, "Модуль предназначен для фиксации запроса и подбора объектов. Сделки, воронка продаж, финансовый учёт, телефония и отдельная система задач в стоимость не входят")

    add_heading(doc, "4 База объектов заявки и каталог новостроек стоимостью 2 500 долларов")
    add_body(doc, "Комплектация включает базу объектов, заявки клиентов и отдельный каталог новостроек")

    add_heading(doc, "4.1 Структура каталога", 2)
    add_body(doc, "В каталоге создаются застройщики, жилые комплексы, корпуса, дома и секции. Для каждого жилого комплекса можно хранить адрес, район, статус строительства, срок сдачи, характеристики, описание, фотографии, документы, внутренние комментарии и контакты")
    add_body(doc, "Жилые комплексы можно добавлять, редактировать, архивировать, искать и фильтровать")

    add_heading(doc, "4.2 Объекты в жилом комплексе", 2)
    add_body(doc, "С жилым комплексом можно связать квартиры и другие предложения. Доступные варианты добавляются вручную или загружаются из подготовленной таблицы")
    add_body(doc, "В карточке жилого комплекса отображается количество связанных предложений. Каждый объект сохраняет собственную цену, площадь, количество комнат, этаж, статус и другие характеристики")
    add_body(doc, "Связанные объекты можно использовать в заявках и клиентских подборках")

    add_heading(doc, "4.3 Наполнение каталога", 2)
    add_body(doc, "Стоимость включает разработку структуры каталога, экранов, полей, связей, поиска, фильтров и инструментов для добавления информации")
    add_body(doc, "Наполнение каталога фотографиями, текстами, документами, конкретными квартирами и актуальными шахматками жилых комплексов выполняется агентством и в стоимость разработки не входит")
    add_body(doc, "Автоматическое получение остатков из систем застройщиков и обработка таблиц произвольного формата согласуются отдельно")

    add_heading(doc, "5 Сравнение комплектаций")
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.55), Inches(1.24), Inches(1.24), Inches(1.24)]
    headers = ["Функционал", "1 800 долларов", "2 000 долларов", "2 500 долларов"]
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "24364B")
        set_cell_margins(cell)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(headers[index])
        configure_font(run, 9.2, True, "FFFFFF")
    set_repeat_table_header(table.rows[0])

    rows = [
        ("База объектов", "входит", "входит", "входит"),
        ("Роли и история изменений", "входит", "входит", "входит"),
        ("Клиентские подборки", "входит", "входит", "входит"),
        ("OLX и DIM RIA", "входит", "входит", "входит"),
        ("Заявки клиентов", "не входит", "входит", "входит"),
        ("Подбор объектов по заявке", "не входит", "входит", "входит"),
        ("Каталог новостроек", "не входит", "не входит", "входит"),
        ("Наполнение каталога новостроек", "не входит", "не входит", "не входит"),
        ("Интеграция с сайтом", "не входит", "не входит", "не входит"),
    ]
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = widths[index]
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            set_cell_border(cells[index])
            if row_index % 2 == 0:
                set_cell_shading(cells[index], "F3F6F8")
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            color = "000000"
            if index > 0 and value == "входит":
                color = "285943"
            elif index > 0 and value == "не входит":
                color = "777777"
            configure_font(run, 9.4, index == 0 or value == "входит", color)

    add_heading(doc, "6 Общие границы проекта")
    add_body(doc, "Во все комплектации входит разработка системы и настройка согласованного функционала")
    add_body(doc, "В текущий состав проекта не входят аренда, посуточная аренда, полноценная CRM, разработка сайта, интеграция с будущим сайтом, серый парсинг объявлений, ручное наполнение базы объектами и постоянное администрирование после запуска")
    add_body(doc, "Интеграции OLX и DIM RIA выполняются при наличии у агентства подходящих корпоративных аккаунтов, разрешений и технического доступа со стороны площадок")
    add_body(doc, "Функции, которые появятся после согласования технического задания, оцениваются отдельно или заменяют другие работы внутри выбранного бюджета по взаимному согласованию")

    heading_7 = add_heading(doc, "7 Срок реализации")
    heading_7.paragraph_format.page_break_before = True
    add_body(doc, "Ориентировочный срок реализации составляет от 6 до 8 недель после согласования технического задания и предоставления необходимых доступов")
    add_body(doc, "Предварительная структура работ включает подготовку технического задания, разработку основной базы, подключение выбранных модулей и интеграций, проверку и запуск")
    add_body(doc, "Срок зависит от выбранной комплектации, скорости согласования и доступности технических кабинетов OLX и DIM RIA")

    add_heading(doc, "8 Переход к техническому заданию")
    add_body(doc, "После выбора комплектации готовится детальное техническое задание. В нём фиксируются все экраны, поля, обязательные данные, права пользователей, статусы, действия, уведомления, интеграции и критерии приёмки")
    add_body(doc, "Окончательная стоимость закрепляется после согласования технического задания в пределах выбранной комплектации")
    add_body(doc, "Порядок оплаты - 20 процентов перед началом подготовки технического задания, 30 процентов после его согласования и 50 процентов после завершения разработки и проверки системы")

    footer = section.footer
    footer.paragraphs[0].text = ""
    add_page_number(footer.paragraphs[0])

    core = doc.core_properties
    core.title = "Предварительное описание проекта"
    core.subject = "Персональная база объектов агентства недвижимости"
    core.author = ""
    core.last_modified_by = ""
    core.keywords = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
